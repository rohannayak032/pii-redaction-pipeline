"""
mock_modules.py - stand-ins for the Extraction and Detection modules until the real ones are connected.
Uses the same shape as the team contract (blocks with file/location/index/text, entities with block-relative
offsets). Everything here is fake-data friendly; nothing is logged or sent anywhere.
"""
import os
import re
from dataclasses import dataclass, field
from typing import Optional

SAMPLE_TEXT = "[MOCK DATA] Prepared by: Jane Smith | Employee ID: EMP-204518 | jane.smith@example.com"


@dataclass
class MockBlock:
    file: str
    location: str
    index: int
    text: str
    page: int = 1
    ocr_confidence: Optional[float] = None


@dataclass
class MockEntity:
    block_index: int
    entity_type: str
    start: int
    end: int
    confidence: float
    detector: str = "regex (mock)"
    value: str = field(default="", repr=False)     # raw PII: never printed

    def public_info(self):
        return {"block_index": self.block_index, "entity_type": self.entity_type, "start": self.start,
                "end": self.end, "confidence": self.confidence, "detector": self.detector}


# ------------------------------------------------------------------ extraction
def extract_document(path, file_type=None):
    path = str(path)
    name = os.path.basename(path)
    ext = (file_type or os.path.splitext(path)[1].lstrip(".")).lower()
    blocks, warns = [], []

    def add(loc, text, page=1):
        if text and text.strip():
            blocks.append(MockBlock(name, loc, len(blocks), text, page))

    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("PDF is password protected")
            for i, pg in enumerate(reader.pages, start=1):
                t = pg.extract_text() or ""
                if not t.strip():
                    warns.append(f"Page {i}: no text layer (scanned? needs OCR from the extraction module).")
                add(f"page {i}", t, i)
        elif ext == "docx":
            import docx
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            d = docx.Document(path)
            for sec in d.sections:
                for p in sec.header.paragraphs:
                    add("header", p.text)
            n, t_i = 0, 0
            for child in d.element.body.iterchildren():
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    n += 1
                    add(f"paragraph {n}", Paragraph(child, d).text)
                elif tag == "tbl":
                    t_i += 1
                    for r_i, row in enumerate(Table(child, d).rows, start=1):
                        add(f"table {t_i}, row {r_i}", " | ".join(c.text.strip() for c in row.cells))
            for sec in d.sections:
                for p in sec.footer.paragraphs:
                    add("footer", p.text)
            warns.append("Mock DOCX extraction skips footnotes and images.")
        elif ext == "pptx":
            from pptx import Presentation
            prs = Presentation(path)

            def walk(shape, slide_no, label):
                if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
                    for s in shape.shapes:
                        walk(s, slide_no, label)
                    return
                if getattr(shape, "has_text_frame", False):
                    add(f"slide {slide_no}, {label}", shape.text_frame.text, slide_no)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for r_i, row in enumerate(shape.table.rows, start=1):
                        add(f"slide {slide_no}, table row {r_i}", " | ".join(c.text.strip() for c in row.cells), slide_no)
            for i, slide in enumerate(prs.slides, start=1):
                for j, sh in enumerate(slide.shapes, start=1):
                    walk(sh, i, f"shape {j}")
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                    add(f"slide {i}, notes", slide.notes_slide.notes_text_frame.text, i)
        elif ext == "txt":
            with open(path, "rb") as f:
                raw = f.read()
            text = None
            for enc in ("utf-8", "utf-16", "latin-1"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeError:
                    continue
            for k, para in enumerate(re.split(r"\n\s*\n", text or ""), start=1):
                add(f"paragraph {k}", para)
        elif ext == "image":
            add("image", SAMPLE_TEXT)
            warns.append("Mock mode has no OCR - dummy text returned for this image.")
        else:
            raise ValueError(f"Mock extractor does not support '{ext}'")
    except ImportError as e:
        add("mock", SAMPLE_TEXT)
        warns.append(f"Library missing for {ext} ({e.name}); dummy text returned.")
    return {"blocks": blocks, "warnings": warns}


# ------------------------------------------------------------------ detection
def _luhn(num):
    digits = [int(c) for c in num if c.isdigit()]
    if not 13 <= len(digits) <= 19:
        return False
    s = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2 == 1:
            d *= 2
            d -= 9 if d > 9 else 0
        s += d
    return s % 10 == 0


_PATTERNS = [  # (type, regex, group, confidence)
    ("EMAIL", r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b", 0, 0.99),
    ("SSN", r"\b\d{3}-\d{2}-\d{4}\b", 0, 0.95),
    ("PHONE", r"(?<![\w])(?:\+?\d{1,3}[\s.\-]?)?(?:\(\d{3}\)|\d{3})[\s.\-]?\d{3}[\s.\-]?\d{4}(?![\w])", 0, 0.85),
    ("CREDIT_CARD", r"\b(?:\d[ \-]?){13,19}\b", 0, 0.97),
    ("IBAN", r"\b[A-Z]{2}\d{2}[A-Z0-9]{11,30}\b", 0, 0.90),
    ("IP_ADDRESS", r"\b(?:(?:25[0-5]|2[0-4]\d|1?\d?\d)\.){3}(?:25[0-5]|2[0-4]\d|1?\d?\d)\b", 0, 0.90),
    ("DATE_OF_BIRTH", r"(?i:\bDOB|date of birth|born(?: on)?)\s*[:\-]?\s*(\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|[A-Z][a-z]+ \d{1,2},? \d{4})", 1, 0.90),
    ("EMPLOYEE_ID", r"\b(?:EMP|EID)[-_ ]?\d{3,8}\b", 0, 0.90),
    ("EMPLOYEE_ID", r"(?i:employee\s*(?:id|number|no\.?))\s*[:#\-]?\s*([A-Za-z0-9\-]{4,12})", 1, 0.85),
    ("PASSPORT", r"(?i:passport(?:\s*(?:no\.?|number|#))?)\s*[:#\-]?\s*([A-Z]{1,2}\d{6,8})\b", 1, 0.85),
    ("PERSON", r"\b(?:Mr|Mrs|Ms|Dr|Prof)\.?[ \t]+[A-Z][a-z]+(?:[ \t]+[A-Z][a-z]+)?", 0, 0.80),
    ("PERSON", r"(?i:name|prepared by|approved by|reviewed by|author|owner|contact|signed by)[ \t]*[:\-][ \t]*([A-Z][a-z]+(?:[ \t]+[A-Z][a-z'\-]+){1,2})", 1, 0.75),
    ("ADDRESS", r"\b\d{1,5}[ \t]+[A-Z][A-Za-z]+(?:[ \t]+[A-Z][A-Za-z]+)*[ \t]+(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|Lane|Ln|Drive|Dr)\b\.?", 0, 0.70),
]
_COMPILED = [(t, re.compile(p), g, c) for t, p, g, c in _PATTERNS]


def detect_pii(blocks):
    out = []
    for pos, b in enumerate(blocks):
        text = b if isinstance(b, str) else (b.get("text") if isinstance(b, dict) else getattr(b, "text", ""))
        bidx = getattr(b, "index", pos) if not isinstance(b, (str, dict)) else pos
        cands = []
        for etype, rx, grp, conf in _COMPILED:
            for m in rx.finditer(text or ""):
                s, e = m.start(grp), m.end(grp)
                val = text[s:e]
                e -= len(val) - len(val.rstrip())
                val = val.rstrip()
                if not val or (etype == "CREDIT_CARD" and not _luhn(val)):
                    continue
                cands.append(MockEntity(bidx, etype, s, e, conf, value=val))
        cands.sort(key=lambda c: (-(c.end - c.start), -c.confidence))
        out.extend(sorted(_dedupe(cands), key=lambda c: c.start))
    return out


def _dedupe(cands):
    kept = []
    for c in cands:        # cands already sorted longest-first
        if all(c.end <= k.start or c.start >= k.end for k in kept):
            kept.append(c)
    return kept
