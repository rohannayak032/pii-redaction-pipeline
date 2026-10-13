"""Generate small fake evidence files (all data is fictional) for testing the demo."""
import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "sample_files"
os.makedirs(OUT, exist_ok=True)

LINES = [
    "Access Control Review - Evidence Record",
    "Prepared by: Jane Smith",
    "Employee ID: EMP-204518",
    "Contact: jane.smith@cadence-example.com | (212) 555-0147",
    "SSN: 123-45-6789",
    "DOB: 04/12/1988",
    "Corporate card: 4111 1111 1111 1111",
    "Approved by Dr. Robert Allen, 42 Maple Street.",
    "Repeat contact: jane.smith@cadence-example.com",
]

# DOCX
import docx
d = docx.Document()
d.sections[0].header.paragraphs[0].text = "Confidential - Owner: Jane Smith"
for l in LINES:
    d.add_paragraph(l)
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Email"
t.cell(1, 0).text, t.cell(1, 1).text = "Priya Raman", "priya.raman@cadence-example.com"
d.save(f"{OUT}/process_doc.docx")

# PPTX
from pptx import Presentation
from pptx.util import Inches
p = Presentation()
for i in range(3):
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = f"Policy Overview {i+1}"
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    tb.text_frame.text = LINES[i + 1] if i < 2 else "No personal data on this slide."
    if i == 1:
        s.notes_slide.notes_text_frame.text = "Call SSN holder: 987-65-4321"
p.save(f"{OUT}/org_deck.pptx")

# PDF
from fpdf import FPDF
pdf = FPDF()
for pg in range(2):
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for l in (LINES if pg == 0 else ["Page two", "Passport: K1234567", "Server IP 192.168.10.44"]):
        pdf.cell(0, 8, l, new_x="LMARGIN", new_y="NEXT")
pdf.output(f"{OUT}/policy_scan.pdf")

# Bad files for error-handling demo
open(f"{OUT}/empty.txt", "wb").close()
open(f"{OUT}/fake.pdf", "wb").write(b"this is not a pdf")
open(f"{OUT}/notes.xyz", "wb").write(b"unsupported")
print("Sample files written to", OUT)
