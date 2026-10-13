"""Build a synthetic PPTX that exercises every place text can hide."""

from __future__ import annotations

import copy
import io
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.packuri import PackURI
from pptx.opc.package import Part
from pptx.util import Emu, Inches, Pt

from . import fake as f

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
RT_LEGACY = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
RT_LEGACY_AUTHORS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors"
RT_MODERN = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
RT_MODERN_AUTHORS = "http://schemas.microsoft.com/office/2018/10/relationships/authors"

# Values only used by the PPTX fixture.
PERSON_GROUP_LEFT = "Ada Lin"
PERSON_GROUP_RIGHT = "Bo Chen"
PERSON_NESTED_GROUP = "Cy Park"
PERSON_NOTES = "Dee Ross"
PERSON_HIDDEN = "Eve Hart"
PERSON_FOOTER = "Flo Banks"
PERSON_MASTER = "Gil Moss"
PERSON_LEGACY_AUTHOR = "Hal Ward"
PERSON_LEGACY_COMMENT = "Ira Cole"
PERSON_MODERN_AUTHOR = "Jo Vance"
PERSON_MODERN_COMMENT = "Kai Ford"
PERSON_CHART = "Liv Grant"
PERSON_UNCLAIMED = "Mo Reyes"
PERSON_TABLE = "Ned Shaw"


def _part(package, partname: str, content_type: str, xml: str) -> Part:
    return Part(PackURI(partname), content_type, package, xml.encode("utf-8"))


def _footer_placeholder(slide, text: str) -> None:
    sp = etree.fromstring(f"""
<p:sp xmlns:p="{P}" xmlns:a="{A}">
  <p:nvSpPr><p:cNvPr id="90" name="Footer Placeholder"/><p:cNvSpPr/><p:nvPr><p:ph type="ftr" idx="11"/></p:nvPr></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="3000000" y="6400000"/><a:ext cx="3000000" cy="300000"/></a:xfrm></p:spPr>
  <p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody>
</p:sp>""")
    slide.shapes._spTree.append(sp)


def build_pptx(path: Path) -> Path:
    prs = Presentation()
    prs.core_properties.author = f.PERSON_AUTHOR
    prs.core_properties.last_modified_by = f.PERSON_LAST_MODIFIED
    blank = prs.slide_layouts[6]

    # --- slide 1: title + two columns of grouped shapes (nested group), footer
    s1 = prs.slides.add_slide(prs.slide_layouts[5])          # "Title Only"
    s1.shapes.title.text = f"Team roster for {f.PERSON_HEADING}"
    group = s1.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(1)).text_frame.text = f"Right: {PERSON_GROUP_RIGHT}"
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = f"Left: {PERSON_GROUP_LEFT}"
    inner = group.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = f"Nested: {PERSON_NESTED_GROUP}"
    s1.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.5)).text_frame.text = f"Contact {f.EMAIL_BODY}"
    _footer_placeholder(s1, f"Owner {PERSON_FOOTER}")
    s1.notes_slide.notes_text_frame.text = f"Remind {PERSON_NOTES} about {f.SSN}"

    # --- slide 2: table with header row, picture with text + alt text, chart
    s2 = prs.slides.add_slide(blank)
    shape = s2.shapes.add_table(3, 3, Inches(0.5), Inches(0.5), Inches(8), Inches(1.5))
    table = shape.table
    for c, label in enumerate(["Employee ID", "Name", "Email"]):
        table.cell(0, c).text = label
    for c, value in enumerate([f.EMP_1, PERSON_TABLE, f.EMAIL_BODY]):
        table.cell(1, c).text = value
    table.cell(2, 0).merge(table.cell(2, 1))
    table.cell(2, 0).text = f"Approved by {f.PERSON_MERGED}"
    image = f.text_image([f"Badge holder: {f.PERSON_IMAGE}", f"ID {f.EMP_2}"])
    pic = s2.shapes.add_picture(io.BytesIO(f.png_bytes(image)), Inches(0.5), Inches(2.5), Inches(5))
    pic._element.nvPicPr.cNvPr.set("descr", f"Photo of {f.PERSON_ALT}")
    tiny = s2.shapes.add_picture(io.BytesIO(f.png_bytes(f.text_image(["x"], size=(20, 20)))), Inches(7), Inches(5))
    data = CategoryChartData()
    data.categories = [PERSON_CHART, "Others"]
    data.add_series("Hours", (10, 20))
    chart = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2.5), Inches(3), Inches(2), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Hours by person"

    # --- slide 3: hidden
    s3 = prs.slides.add_slide(blank)
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = f"Hidden: {PERSON_HIDDEN}"
    s3._element.set("show", "0")
    # text in an extension list: only the sweep can find it
    ext = etree.fromstring(
        f'<p:extLst xmlns:p="{P}" xmlns:a="{A}"><p:ext uri="{{00000000-0000-0000-0000-000000000000}}">'
        f'<a:p><a:r><a:t>Unclaimed {PERSON_UNCLAIMED}</a:t></a:r></a:p></p:ext></p:extLst>'
    )
    s3._element.append(ext)

    # --- master: a plain text box shown on every slide
    scratch = prs.slides.add_slide(blank)
    box = scratch.shapes.add_textbox(Inches(0.2), Inches(7), Inches(4), Inches(0.4))
    box.text_frame.text = f"Prepared by {PERSON_MASTER}"
    prs.slide_masters[0].shapes._spTree.append(copy.deepcopy(box._element))
    _delete_last_slide(prs)

    # --- comments: legacy on slide 1, modern on slide 2
    package = prs.part.package
    legacy_authors = _part(package, "/ppt/commentAuthors.xml",
                           "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml",
                           f'<p:cmAuthorLst xmlns:p="{P}"><p:cmAuthor id="1" name="{PERSON_LEGACY_AUTHOR}" initials="HW" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>')
    prs.part.relate_to(legacy_authors, RT_LEGACY_AUTHORS)
    legacy = _part(package, "/ppt/comments/comment1.xml",
                   "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
                   f'<p:cmLst xmlns:p="{P}"><p:cm authorId="1" idx="1"><p:pos x="10" y="10"/>'
                   f'<p:text>Check with {PERSON_LEGACY_COMMENT}</p:text></p:cm></p:cmLst>')
    s1.part.relate_to(legacy, RT_LEGACY)
    modern_authors = _part(package, "/ppt/authors.xml",
                           "application/vnd.ms-powerpoint.authors+xml",
                           f'<p188:authorLst xmlns:p188="{P188}"><p188:author id="{{A1}}" name="{PERSON_MODERN_AUTHOR}" initials="JV" userId="x" providerId="None"/></p188:authorLst>')
    prs.part.relate_to(modern_authors, RT_MODERN_AUTHORS)
    modern = _part(package, "/ppt/comments/modernComment_1.xml",
                   "application/vnd.ms-powerpoint.comments+xml",
                   f'<p188:cmLst xmlns:p188="{P188}" xmlns:a="{A}"><p188:cm id="{{C1}}" authorId="{{A1}}">'
                   f'<p188:txBody><a:bodyPr/><a:p><a:r><a:t>Ping {PERSON_MODERN_COMMENT}</a:t></a:r></a:p></p188:txBody>'
                   '</p188:cm></p188:cmLst>')
    s2.part.relate_to(modern, RT_MODERN)

    prs.save(str(path))
    return path


def _delete_last_slide(prs) -> None:
    sld_ids = prs.slides._sldIdLst
    last = sld_ids[-1]
    prs.part.drop_rel(last.rId)
    sld_ids.remove(last)
