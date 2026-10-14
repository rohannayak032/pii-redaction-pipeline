"""Write data/sample_addresses.docx: a synthetic document with Indian addresses for the demo,
and its ground truth data/evaluation_results/_final_gt/sample_addresses.docx.csv for the app's Evaluate tab.

Every person, house, building and e-mail is invented (e-mails use the reserved example.com domain);
localities, cities, states and PIN codes are real places so the addresses look realistic. No phone
numbers: any 10-digit Indian mobile number could belong to someone.

The addresses come in the shapes real documents use: inside a sentence, in a table cell, as a
multi-line letter block, in upper case with a dashed PIN, as a rural village/tehsil/district
address, with a "PIN -" label and in the page footer. City names on their own are not addresses
and must not be redacted.

    python src/pii_pipeline/app/report/make_address_demo.py           # write the file and its ground truth
    python src/pii_pipeline/app/report/make_address_demo.py --check   # also run extraction + detection on it
"""
import csv
import sys
from pathlib import Path

import docx

DATA = Path(__file__).resolve().parents[4] / "data"
OUT = DATA / "sample_addresses.docx"
GT = DATA / "evaluation_results" / "_final_gt" / "sample_addresses.docx.csv"

# (sentence, person named in it or None, address)
SENTENCES = [
    ("Ananya Raghavan, our new branch manager, will receive all deliveries at {} from Monday.",
     "Ananya Raghavan",
     "Flat 7B, Saffron Heights, 12th Cross, Indiranagar, Bengaluru, Karnataka 560038"),
    # A name right before "at <address>" gets pulled into the address span, so names sit apart.
    ("Vikram Desai has moved, so please send the signed agreement to {} before the end of the month.",
     "Vikram Desai",
     "House No. 221, Gulmohar Lane, Sector 15, Vashi, Navi Mumbai, Maharashtra 400703"),
    ("The Hyderabad warehouse moved to {} and the old site is closed.",
     None,
     "Plot 46, Lotus Enclave, Road No. 3, Banjara Hills, Hyderabad, Telangana 500034"),
    ("Meera Iyer asked that her payslips be posted to {} instead of the office.",
     "Meera Iyer",
     "D-118, Neelkanth Apartments, Mayur Vihar Phase 1, East Delhi, Delhi 110091"),
    ("During the harvest season, leave letters for Rahul Verma are sent to {}.",
     "Rahul Verma",
     "H.No. 56, Village Rampur, Post Office Bhadohi, District Bhadohi, Uttar Pradesh 221401"),
    ("Fatima Sheikh changed her mailing address. New address: {}.",
     "Fatima Sheikh",
     "Door No. 12/345, Chittoor Road, Ernakulam, Kochi, Kerala, PIN - 682016"),
]
# Place names that are not addresses: nothing here may be redacted.
NEGATIVES = [
    "The Pune and Chennai teams meet every quarter; no visit is planned for Kolkata this year.",
    "Shipments through the Mumbai, Delhi and Hyderabad hubs grew 12% this quarter.",
]
LETTER_NAME = "Devendra Joshi"
LETTER_ADDRESS = "Shop No. 4, Ground Floor, Shanti Complex\nStation Road, Navrangpura\nAhmedabad, Gujarat 380009"
LETTER_LINES = ["To,", LETTER_NAME, "Proprietor, Joshi Packaging Works", LETTER_ADDRESS]
TABLE = [
    ("Karthik Subramanian", "Field Engineer", "No. 9, Kamaraj Street, T. Nagar, Chennai, Tamil Nadu 600017",
     "karthik.subramanian@example.com"),
    ("Priya Kulkarni", "Accounts Executive", "B-302, Riverstone Residency, Baner Road, Pune, Maharashtra 411045",
     "priya.kulkarni@example.com"),
    ("Arjun Banerjee", "Logistics Lead", "14/2, Lake View Colony, Salt Lake Sector 2, Kolkata, West Bengal 700091",
     "arjun.banerjee@example.com"),
    ("Sneha Rathore", "HR Partner", "Villa 33, Palm Grove, Vaishali Nagar, Jaipur, Rajasthan 302021",
     "sneha.rathore@example.com"),
    ("Imran Qureshi", "Branch Auditor",
     "FLAT 502, SAI KRUPA APARTMENTS, 2ND MAIN ROAD, JAYANAGAR 4TH BLOCK, BENGALURU, KARNATAKA - 560011",
     "imran.qureshi@example.com"),
]
FOOTER_PREFIX = "Northwind Logistics India Pvt. Ltd. | Registered office: "
FOOTER_ADDRESS = "Unit 501, 5th Floor, Orchid Tower, Golf Course Road, Sector 54, Gurugram, Haryana 122002"
ADDRESSES = [a for _, _, a in SENTENCES] + [LETTER_ADDRESS] + [row[2] for row in TABLE] + [FOOTER_ADDRESS]


def ground_truth() -> list[list[str]]:
    """GT rows in the _final_gt format; values without offsets are matched by file + type + value."""
    f = OUT.name
    rows = []
    for i, (_, name, address) in enumerate(SENTENCES, 1):
        loc = f"Correspondence notes, paragraph {i}"
        if name:
            rows.append([f, "IND-ADDR-NOTES", "NAME", name, loc])
        rows.append([f, "IND-ADDR-NOTES", "ADDRESS", address, loc])
    rows.append([f, "IND-ADDR-LETTER", "NAME", LETTER_NAME, "Vendor letter, recipient block"])
    rows.append([f, "IND-ADDR-LETTER", "ADDRESS", LETTER_ADDRESS, "Vendor letter, recipient block"])
    for i, (name, _, address, email) in enumerate(TABLE, 1):
        loc = f"Employee address register, table, row {i}"
        rows += [[f, "IND-ADDR-TABLE", t, v, loc] for t, v in (("NAME", name), ("ADDRESS", address), ("EMAIL", email))]
    rows.append([f, "IND-ADDR-FOOTER", "ADDRESS", FOOTER_ADDRESS, "page footer"])
    # Line breaks become spaces: the evaluator collapses whitespace before its leak check.
    return [r[:3] + [" ".join(r[3].split()), r[4], "", ""] for r in rows]


def write(path: Path = OUT, gt_path: Path = GT) -> Path:
    d = docx.Document()
    d.sections[0].footer.paragraphs[0].text = FOOTER_PREFIX + FOOTER_ADDRESS
    d.add_heading("Northwind Logistics India - Branch Onboarding Register", 1)
    d.add_paragraph("SYNTHETIC DEMO DATA. Every name, address and e-mail in this document is invented "
                    "to test PII redaction.")
    d.add_heading("Correspondence notes", 2)
    for sentence, _, address in SENTENCES:
        d.add_paragraph(sentence.format(address))
    for sentence in NEGATIVES:
        d.add_paragraph(sentence)
    d.add_heading("Vendor letter", 2)
    d.add_paragraph("\n".join(LETTER_LINES))   # one paragraph, lines split by soft line breaks
    d.add_heading("Employee address register", 2)
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Employee", "Role", "Correspondence address", "E-mail")):
        cell.text = text
    for row in TABLE:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)

    gt_path.parent.mkdir(parents=True, exist_ok=True)
    with open(gt_path, "w", newline="", encoding="utf-8") as o:
        csv.writer(o).writerows([["file", "case_id", "type", "value", "location", "start", "end"]] + ground_truth())
    return path


def check(path: Path) -> None:
    """Every address must sit inside one ADDRESS detection (the redaction then removes all of it),
    then the whole document is scored against its ground truth."""
    from pii_pipeline.detection.detector import detect
    from pii_pipeline.evaluation.evaluate_redaction import evaluate_redaction, format_report
    from pii_pipeline.extraction import extract_document
    from pii_pipeline.redaction import redact_document

    blocks = extract_document(path)
    entities = detect(blocks)
    missed = 0
    for address in ADDRESSES:
        address = " ".join(address.split())   # extraction turns the letter's line breaks into spaces
        block = next(b for b in blocks if address in b.text)
        start = block.text.index(address)
        covered = any(e.pii_type == "ADDRESS" and e.block_index == block.index
                      and e.start <= start and start + len(address) <= e.end for e in entities)
        missed += not covered
        print(f"{'ok  ' if covered else 'MISS'} {address}")
    by_detector = {d: sum(e.detector == d for e in entities) for d in sorted({e.detector for e in entities})}
    print(f"{len(ADDRESSES) - missed}/{len(ADDRESSES)} addresses fully covered; detections by detector: {by_detector}")
    redacted, _ = redact_document(blocks, entities)
    print(format_report(evaluate_redaction(str(GT), entities, redacted)))
    assert missed == 0


if __name__ == "__main__":
    out = write()
    print(f"wrote {out}\nwrote {GT}")
    if "--check" in sys.argv:
        check(out)
