"""PPTX extraction.

Per slide (hidden slides included and marked ``(hidden)``): shapes in visual
order, then speaker notes (and other notes-page text), then comments. After
all slides: master/layout text, the safety-net sweep, document properties.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from ..base import Context, RawBlock
from ..metadata import ooxml_properties
from ..ooxml import Consumed, count_embedded, drawingml_text, part_name, sweep
from ..validation import UNSUPPORTED_OBJECT
from .comments import author_names, comment_blocks
from .masters import template_blocks
from .shapes import SlideWalker

def _notes_blocks(slide, consumed: Consumed, prefix: str, page: int, hidden: bool) -> list[RawBlock]:
    if not slide.has_notes_slide:
        return []
    notes = slide.notes_slide
    name = part_name(notes.part)
    blocks: list[RawBlock] = []
    body = notes.notes_placeholder
    if body is not None:
        text = drawingml_text(body._element, name, consumed)
        if text:
            blocks.append(RawBlock.simple("notes", f"{prefix}, speaker notes", text, source="xml", page=page, hidden=hidden))
    for shape in notes.shapes:
        if body is not None and shape.shape_id == body.shape_id:
            continue
        text = drawingml_text(shape._element, name, consumed)
        if not text:
            continue
        if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        blocks.append(RawBlock.simple(
            "notes", f"{prefix}, notes page, shape {shape.shape_id}", text, source="xml", page=page, hidden=hidden,
        ))
    return blocks


def extract_pptx(path: Path, ctx: Context) -> list[RawBlock]:
    prs = Presentation(str(path))
    consumed = Consumed()
    authors = author_names(prs.part)
    blocks: list[RawBlock] = []

    for number, slide in enumerate(prs.slides, start=1):
        hidden = slide._element.get("show") in ("0", "false")
        prefix = f"slide {number}" + (" (hidden)" if hidden else "")
        walker = SlideWalker(ctx, consumed, slide.part, prefix, number, hidden)
        blocks += walker.walk(slide.shapes)
        blocks += _notes_blocks(slide, consumed, prefix, number, hidden)
        blocks += comment_blocks(slide.part, authors, consumed, prefix, number, hidden)

    blocks += template_blocks(prs, consumed)

    package = prs.part.package
    embedded = count_embedded(package)
    if embedded:
        ctx.warn(UNSUPPORTED_OBJECT, "document", f"{embedded} embedded object(s)/file(s) not extracted")

    blocks += sweep(package.iter_parts(), consumed, ctx, prefixes=("ppt/",))
    properties = ooxml_properties(package)
    if properties is not None:
        blocks.append(properties)
    return blocks
