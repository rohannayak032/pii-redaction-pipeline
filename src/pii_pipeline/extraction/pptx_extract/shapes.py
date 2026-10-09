"""Slide shapes -> blocks, in visual reading order.

* Groups are walked recursively; child coordinates are converted to slide
  coordinates through each group's transform, and a group's children stay
  together (ordered among themselves by geometry).
* Top level: title placeholders first, then XY-cut over everything else,
  then footer/date placeholders. Slide-number placeholders carry no content.
* Locations use shape ids, never shape names: names are free text an author
  can set to anything, including a person's name.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.picture import Picture

from ..base import BBox, Context, Piece, RawBlock
from ..normalize import clean
from ..ocr import ocr_blob
from ..ooxml import Consumed, DIAGRAM_URI, NS, chart_text, drawingml_text, part_name, smartart_text, xp
from ..ordering import union, xy_cut
from ..tables import Cell, table_blocks
from ..validation import UNSUPPORTED_OBJECT

_TITLE = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_FOOTER = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}
_OLE_URI = "http://schemas.openxmlformats.org/presentationml/2006/ole"
_XFRM = xp("./p:grpSpPr/a:xfrm")
_CNVPR = xp("./*[1]/p:cNvPr")
_GRAPHIC_DATA = xp("./a:graphic/a:graphicData")
_DGM_RELIDS = xp("./a:graphic/a:graphicData/dgm:relIds")
_TABLE = xp("./a:graphic/a:graphicData/a:tbl")
_R = f"{{{NS['r']}}}"


@dataclass(frozen=True)
class Transform:
    """Child -> slide coordinates: x' = x * sx + tx (same for y)."""

    sx: float = 1.0
    sy: float = 1.0
    tx: float = 0.0
    ty: float = 0.0

    def apply(self, x: float, y: float) -> tuple[float, float]:
        return x * self.sx + self.tx, y * self.sy + self.ty

    def child(self, group) -> "Transform":
        xfrm = next(iter(_XFRM(group._element)), None)
        if xfrm is None:
            return self
        def pair(tag: str, a: str, b: str) -> tuple[float, float]:
            node = xfrm.find(f"{{{NS['a']}}}{tag}")
            return (float(node.get(a, 0)), float(node.get(b, 0))) if node is not None else (0.0, 0.0)
        ox, oy = pair("off", "x", "y")
        cx, cy = pair("ext", "cx", "cy")
        chx, chy = pair("chOff", "x", "y")
        chcx, chcy = pair("chExt", "cx", "cy")
        kx = cx / chcx if chcx else 1.0
        ky = cy / chcy if chcy else 1.0
        return Transform(
            sx=self.sx * kx, sy=self.sy * ky,
            tx=self.sx * (ox - chx * kx) + self.tx,
            ty=self.sy * (oy - chy * ky) + self.ty,
        )


@dataclass
class Item:
    bbox: BBox
    rank: int                       # 0 title, 1 content, 2 footer
    blocks: list[RawBlock] = field(default_factory=list)


class SlideWalker:
    def __init__(self, ctx: Context, consumed: Consumed, part, prefix: str, page: int, hidden: bool):
        self.ctx, self.consumed, self.part = ctx, consumed, part
        self.name = part_name(part)
        self.prefix, self.page, self.hidden = prefix, page, hidden
        self.tables = self.charts = self.diagrams = self.images = 0

    def _common(self) -> dict:
        return dict(source="xml", page=self.page, hidden=self.hidden)

    def _bbox(self, shape, transform: Transform) -> BBox:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        if left is None or top is None:
            return (0.0, 0.0, 0.0, 0.0)
        x0, y0 = transform.apply(left, top)
        x1, y1 = transform.apply(left + (width or 0), top + (height or 0))
        return (x0, y0, x1, y1)

    def _rank(self, shape) -> int:
        if not shape.is_placeholder:
            return 1
        kind = shape.placeholder_format.type
        return 0 if kind in _TITLE else 2 if kind in _FOOTER else 1

    # ------------------------------------------------------------------ shapes
    def shape_item(self, shape, transform: Transform) -> Item | None:
        bbox = self._bbox(shape, transform)
        rank = self._rank(shape)
        blocks: list[RawBlock] = []
        sid = shape.shape_id

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            inner = [it for it in (self.shape_item(s, transform.child(shape)) for s in shape.shapes) if it]
            for i in xy_cut([it.bbox for it in inner]):
                blocks += inner[i].blocks
        elif shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            drawingml_text(shape._element, self.name, self.consumed)   # mark consumed, no content
        elif getattr(shape, "has_text_frame", False):
            text = drawingml_text(shape._element, self.name, self.consumed)
            if text:
                where = {0: "title", 2: "footer"}.get(rank, f"shape {sid}")
                kind = {0: "title", 2: "footer"}.get(rank, "text")
                blocks.append(RawBlock.simple(kind, f"{self.prefix}, {where}", text, bbox=bbox, bbox_unit="emu", **self._common()))
        elif getattr(shape, "has_table", False):
            blocks += self._table(shape, bbox)
        elif getattr(shape, "has_chart", False):
            text = chart_text(shape.chart_part, self.consumed)
            if text:
                self.charts += 1
                blocks.append(RawBlock.simple("chart", f"{self.prefix}, chart {self.charts}", text, bbox=bbox, bbox_unit="emu", **self._common()))
        elif _DGM_RELIDS(shape._element):
            rid = _DGM_RELIDS(shape._element)[0].get(f"{_R}dm")
            text = smartart_text(self.part, rid, self.consumed) if rid else ""
            if text:
                self.diagrams += 1
                blocks.append(RawBlock.simple("smartart", f"{self.prefix}, diagram {self.diagrams}", text, bbox=bbox, bbox_unit="emu", **self._common()))
        elif any(gd.get("uri") == _OLE_URI for gd in _GRAPHIC_DATA(shape._element)):
            self.ctx.warn(UNSUPPORTED_OBJECT, f"{self.prefix}, shape {sid}", "embedded object (OLE) content not extracted")

        if isinstance(shape, Picture):
            blocks += self._picture(shape)
        blocks += self._alt_text(shape, bbox)
        return Item(bbox, rank, blocks) if blocks else None

    def _picture(self, shape) -> list[RawBlock]:
        try:
            blob = shape.image.blob
        except (ValueError, AttributeError, KeyError):
            self.ctx.skip_image("linked (external) or missing image")
            return []
        self.images += 1
        blocks = ocr_blob(blob, self.ctx, f"{self.prefix}, image {self.images}", hidden=self.hidden)
        for block in blocks:
            block.page = self.page
        return blocks

    def _alt_text(self, shape, bbox: BBox) -> list[RawBlock]:
        node = next(iter(_CNVPR(shape._element)), None)
        if node is None:
            return []
        alt = " ".join(v for v in (node.get("title"), node.get("descr")) if clean(v))
        if not alt:
            return []
        return [RawBlock(
            kind="alt_text", location=f"{self.prefix}, shape {shape.shape_id} alt text",
            pieces=[Piece("Alt text", kind="label"), Piece(": ", kind="sep"), Piece(alt)],
            bbox=bbox, bbox_unit="emu", **self._common(),
        )]

    def _table(self, shape, bbox: BBox) -> list[RawBlock]:
        tbl = _TABLE(shape._element)[0]
        rows: list[list[Cell]] = []
        for tr in tbl.iterchildren(f"{{{NS['a']}}}tr"):
            cells: list[Cell] = []
            for col, tc in enumerate(tr.iterchildren(f"{{{NS['a']}}}tc")):
                if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
                    continue
                cells.append(Cell(col=col, text=drawingml_text(tc, self.name, self.consumed)))
            rows.append(cells)
        self.tables += 1
        first_row = shape.table.first_row
        return table_blocks(
            rows, location=f"{self.prefix}, table {self.tables}", table_id=self.ctx.next_table_id(),
            flagged_header_rows=1 if first_row else 0, bbox=bbox, bbox_unit="emu", **self._common(),
        )

    # ------------------------------------------------------------------ slide
    def walk(self, shapes) -> list[RawBlock]:
        items = [it for it in (self.shape_item(s, Transform()) for s in shapes) if it]
        out: list[RawBlock] = []
        for rank in (0, 1, 2):
            group = [it for it in items if it.rank == rank]
            for i in xy_cut([it.bbox for it in group]):
                out += group[i].blocks
        return out
